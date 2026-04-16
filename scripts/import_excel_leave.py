"""Import legacy leave data from the 'Vacation SE' sheet of the Excel file.

Reads temp/Vacation and time off.xlsx and turns each contiguous coloured run
into a LeaveRequest document in Cosmos DB.

Usage (from repo root):
    .venv\\Scripts\\python scripts/import_excel_leave.py [--dry-run] [--year YYYY]

Flags:
    --dry-run        Print what would be imported without writing to Cosmos.
    --year YYYY      Only import leave starting in this year (e.g. 2026).
                     Omit to import everything.
"""

import argparse
import asyncio
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent / "backend"))

from app.db.cosmos import (
    list_items,
    upsert_item,
    USERS_CONTAINER,
    LEAVE_REQUESTS_CONTAINER,
)

# ─── Config ──────────────────────────────────────────────────────────────────

EXCEL_PATH = Path(__file__).parent.parent / "temp" / "Vacation and time off.xlsx"
SHEET_NAME = "Vacation SE"

# Column E (0-based index 4) is the first date column.
# That cell represents Monday 30 Dec 2024 (= ISO week-1 Monday of 2025).
DATA_COL_OFFSET = 4
START_DATE = date(2024, 12, 30)

# Person rows begin at row 7 (1-indexed); rows 1–6 are headers / legend.
DATA_ROW_START = 7

# ─── Excel code → (leaveType, status) ────────────────────────────────────────
# leaveType: "A" = vacation, "FL" = parental, "C" = comp time
# status:    "A" = pending,  "B" = approved
CODE_MAP: dict[str, tuple[str, str]] = {
    "A":  ("A",  "A"),   # Requested vacation
    "B":  ("A",  "B"),   # Approved vacation
    "D":  ("A",  "B"),   # Also approved (legacy "D day" code)
    "FL": ("FL", "B"),   # Parental leave (föräldraledighet)
    "Fl": ("FL", "B"),   # Case variant of FL
    "PL": ("FL", "B"),   # Alternate parental code
    "C":  ("C",  "B"),   # Comp time (fridity/overtime compensation)
}

# Codes to silently skip (sick leave, whitespace noise, unknowns)
SKIP_CODES = {"SJ", "Sj", "?"}


# ─── Helpers ─────────────────────────────────────────────────────────────────

def clean_name(raw: str) -> str:
    """Strip employee IDs like '(40114597)' and whitespace from a display name."""
    return re.sub(r"\s*\(.*?\)", "", raw).strip()


def slugify(text: str) -> str:
    """Convert a name to a lowercase hyphen-slug (matches seed_db.py logic)."""
    text = text.lower()
    for src, dst in [
        ("ä", "a"), ("å", "a"), ("ö", "o"), ("é", "e"), ("ę", "e"),
        ("ó", "o"), ("ń", "n"), ("ł", "l"), ("ź", "z"), ("ż", "z"),
        ("ą", "a"), ("ś", "s"), ("ć", "c"), ("ü", "u"),
    ]:
        text = text.replace(src, dst)
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def user_id_from_name(name: str) -> str:
    return f"user-{slugify(name)}"


def bu_id_from_name(bu_name: str) -> str:
    return f"bu-{slugify(bu_name)}"


def col_to_date(col_0based: int) -> date:
    """Map a 0-based column index to the corresponding calendar date."""
    return START_DATE + timedelta(days=col_0based - DATA_COL_OFFSET)


def extract_runs(indexed_cells: list[tuple[int, object]]):
    """
    Yield (start_col, end_col, excel_code) for each contiguous run of the
    same mappable code in the indexed_cells list.
    Gaps (None / whitespace / skip codes) break runs.
    """
    run_start: int | None = None
    run_code: str | None = None

    for col, raw in indexed_cells:
        code = str(raw).strip() if raw is not None else ""
        is_valid = code in CODE_MAP

        if not is_valid:
            # Close any open run
            if run_start is not None:
                yield (run_start, col - 1, run_code)
                run_start = None
                run_code = None
        else:
            if code != run_code:
                # Close previous run if the code changes
                if run_start is not None:
                    yield (run_start, col - 1, run_code)
                run_start = col
                run_code = code

    # Close final run
    if run_start is not None and indexed_cells:
        yield (run_start, indexed_cells[-1][0], run_code)


# ─── Main import logic ────────────────────────────────────────────────────────

async def do_import(dry_run: bool, year_filter: int | None) -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    max_col_0based = ws.max_column - 1  # convert to 0-based

    first_date = col_to_date(DATA_COL_OFFSET)
    last_date  = col_to_date(max_col_0based)
    print(f"Spreadsheet date range: {first_date} → {last_date}")
    if year_filter:
        print(f"Filtering to year: {year_filter}")
    print()

    # Load existing Cosmos users so we can warn on unknown names
    cosmos_users = {u["id"]: u for u in await list_items(USERS_CONTAINER)}
    print(f"{len(cosmos_users)} users found in Cosmos.\n")

    total_requests = 0
    warned_users: list[str] = []

    for row in ws.iter_rows(min_row=DATA_ROW_START, max_row=ws.max_row, values_only=True):
        raw_name = row[0]
        if not isinstance(raw_name, str) or not raw_name.strip():
            continue

        name    = clean_name(raw_name)
        bu_raw  = row[2]
        if not isinstance(bu_raw, str) or not bu_raw.strip():
            continue
        bu_raw = bu_raw.strip()

        uid     = user_id_from_name(name)
        bu_cosm = bu_id_from_name(bu_raw)

        if uid not in cosmos_users:
            warned_users.append(f"  [WARN] Not in Cosmos: {name!r} → tried id={uid!r}")
            continue

        # Build the indexed cell list for this person's date columns
        indexed_cells = [
            (col, row[col] if col < len(row) else None)
            for col in range(DATA_COL_OFFSET, max_col_0based + 1)
        ]

        person_count = 0
        for start_col, end_col, excel_code in extract_runs(indexed_cells):
            leave_type, status = CODE_MAP[excel_code]
            start_dt = col_to_date(start_col)
            end_dt   = col_to_date(end_col)

            if year_filter and start_dt.year != year_filter:
                continue

            # Deterministic ID: allows safe re-runs (upsert = idempotent)
            lr_id = f"lr-{uid}-{start_dt.isoformat()}-{leave_type.lower()}"

            doc = {
                "id":             lr_id,
                "userId":         uid,
                "businessUnitId": bu_cosm,
                "startDate":      start_dt.isoformat(),
                "endDate":        end_dt.isoformat(),
                "leaveType":      leave_type,
                "status":         status,
                "notes":          f"Imported from Excel (code: {excel_code})",
            }

            if dry_run:
                print(f"  [DRY] {name:35s} {start_dt} → {end_dt}  type={leave_type} status={status}  ({excel_code})")
            else:
                await upsert_item(LEAVE_REQUESTS_CONTAINER, doc)

            person_count += 1
            total_requests += 1

        if person_count > 0 and not dry_run:
            print(f"  ✓ {name}: {person_count} request(s) imported")

    print()
    if warned_users:
        print("Warnings (check names match seed):")
        for w in warned_users:
            print(w)
        print()

    action = "would be imported" if dry_run else "imported"
    print(f"Done. {total_requests} leave requests {action}.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true",
                        help="Print what would be imported without writing to Cosmos")
    parser.add_argument("--year", type=int, default=None,
                        help="Only import leave starting in this year (e.g. 2026)")
    args = parser.parse_args()
    asyncio.run(do_import(dry_run=args.dry_run, year_filter=args.year))
